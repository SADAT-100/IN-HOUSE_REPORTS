"""Turn a Noon KSA workbook (courier 12539 or bulky 18947) into the compact data.json the Noon portal embeds.
Usage: python3 -I prepare_data.py <workbook.xlsx> <out.json>"""
import sys, json, re, datetime as dt
import openpyxl

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

def num(v):
    if v is None or isinstance(v, str) and not re.match(r'^\s*-?[\d.,]+\s*$', v): return 0.0
    try: return float(str(v).replace(',', '')) if isinstance(v, str) else float(v)
    except Exception: return 0.0
r2 = lambda x: round(x + 0.0, 2)
BASE = dt.datetime(2026, 1, 1)
mins = lambda d: int((d - BASE).total_seconds() // 60) if isinstance(d, dt.datetime) else None

# ---- product titles from the statement ----
titles = {}
PS = next(n for n in wb.sheetnames if n.startswith('PS-'))
ACCT = PS[3:]
BULKY = ACCT == '18947'
ws = wb[PS]
rows = ws.iter_rows(values_only=True); H = list(next(rows)); ix = {h: i for i, h in enumerate(H)}
stm = {}
for r in rows:
    if r[ix['Reference Nr']] is None: continue
    sku = r[ix['SKUs']]
    if sku and r[ix['Title']] and sku not in titles: titles[sku] = str(r[ix['Title']]).strip()
    ref = r[ix['Reference Nr']]; ttype = r[ix['Transaction Type']]
    s = stm.setdefault(ref, {'ref': ref, 'date': r[ix['Transaction Date']], 'rows': 0, 'np': 0, 'referral': 0, 'fulfil': 0, 'shipCredit': 0, 'otherOrder': 0, 'subsidy': 0, 'nonOrderFee': 0, 'nonOrderSubsidy': 0, 'others': 0, 'total': 0, 'orderTotal': 0, 'updateTotal': 0})
    s['rows'] += 1
    for k, col in [('np', 'Net Proceeds'), ('referral', 'Referral Fee including VAT'), ('fulfil', 'Fullfilment & Logistics Fees including VAT'), ('shipCredit', 'Shipping Credits including VAT'), ('otherOrder', 'Other Order Fees including VAT'), ('subsidy', 'Order Subsidies including VAT'), ('nonOrderFee', 'Non-Order Fees including VAT'), ('nonOrderSubsidy', 'Non-Order Subsidies including VAT'), ('others', 'Others including VAT'), ('total', 'Total')]:
        s[k] += num(r[ix[col]])
    if ttype == 'order': s['orderTotal'] += num(r[ix['Total']])
    elif ttype == 'order_update': s['updateTotal'] += num(r[ix['Total']])
statements = []
for s in sorted(stm.values(), key=lambda s: s['date']):
    s['date'] = s['date'].strftime('%Y-%m-%d')
    statements.append({k: (r2(v) if isinstance(v, float) else v) for k, v in s.items()})

# ---- items ----
ws = wb[next(n for n in wb.sheetnames if n.startswith('sales_export'))]
rows = ws.iter_rows(values_only=True); H = list(next(rows))
seen = {}; cols = []
for h in H: seen[h] = seen.get(h, 0) + 1; cols.append(h if seen[h] == 1 else f'{h}#2')
ix = {h: i for i, h in enumerate(cols)}
STATUS = ['Delivered', 'Shipped', 'Processing', 'Pending Confirmation', 'CIR', 'Could Not Be Delivered', 'Cancelled']
FUL = {'Fulfilled by Partner (FBP)': 'Flex', 'Fulfilled by Noon (FBN)': 'FBN', 'Supermall': 'Supermall'}
CH = ['Flex', 'FBN', 'Supermall']
if BULKY:  # bulky splits by the delivery type column instead of the fulfilment model
    SHIP = {'Flex_Bulky': 'Flex', 'Dirc_Bulky': 'Direct', 'Noon': 'FBN'}
    CH = ['Flex', 'Direct', 'New Ansari', 'FBN', 'Not set']
    def chan(g):
        t = g('Type Ship')
        if t in SHIP: return SHIP[t]
        if t and str(t).startswith('New_Ansari'): return 'New Ansari'
        return 'FBN' if g('fulfillment_model') == 'Fulfilled by Noon (FBN)' else 'Not set'
else:
    chan = lambda g: FUL[g('fulfillment_model')]
SOLD = {'Delivered', 'Shipped', 'Processing', 'Pending Confirmation'}
def inhouse_fees(g):
    if not BULKY: return [r2(num(g('Comm'))), r2(num(g('Shipping'))), r2(num(g('Order Subsidies & VAT#2')))]
    # Bulky keeps In House commission and shipping as positive costs and has no subsidy column.
    if g('status') == 'Cancelled': return [0.0, 0.0, 0.0]
    if g('status') in SOLD: return [r2(-num(g('Comm'))), r2(-num(g('Shipping'))), 0.0]
    return [0.0, r2(num(g('Revenue2'))), 0.0]  # returned items only carry the return shipping
skus, skuIx, items = [], {}, []
for r in rows:
    g = lambda c: r[ix[c]]
    if not g('item_nr'): continue
    sku = str(g('sku'))
    if sku not in skuIx:
        skuIx[sku] = len(skus)
        ps = g('partner_sku'); ps = ('%d' % ps) if isinstance(ps, float) and ps.is_integer() else str(ps if ps is not None else '')
        skus.append([sku, ps, titles.get(sku, ''), g('Cat') or '', g('family') or '', g('brand_code') or ''])
    order = str(g('order_nr') if 'order_nr' in ix else g('item')); suffix = str(g('item_nr'))[len(order) + 1:]
    items.append([mins(g('order_timestamp')), order, suffix, skuIx[sku], STATUS.index(g('status')), CH.index(chan(g)),
        r2(num(g('gmv_lcy'))), r2(num(g('Net Proceeds'))), r2(num(g('Referal  & Vat'))), r2(num(g('Shipping cost & Vat'))), r2(num(g('Order Subsidies & VAT'))), r2(num(g('Cost'))),
        *inhouse_fees(g),
        r2(num(g('Revenue'))), r2(num(g('Profit'))), r2(num(g('Revenue2'))), r2(num(g('Profit2'))),
        mins(g('shipment_timestamp')), mins(g('delivered_timestamp'))])

# ---- statement & payment sheet ----
ws = wb['Statment&Payment']
month, vat, payments, fbn, stLines = 0, [], [], [], []
SP = list(ws.iter_rows(values_only=True)); PD = list(SP[0]).index('Payment')
def pay_date(s, m):
    a, b, y = [int(x) for x in s.split('.')]
    opts = [(b, a), (a, b)]  # (month, day): dd.mm first, then mm.dd
    for mm, dd in opts:
        if mm == m and dd <= 31: return dt.date(y, mm, dd)
    for mm, dd in opts:
        if mm == m + 1 and dd <= 31: return dt.date(y, mm, dd)
    return None
for r in SP[1:]:
    if r[0] and str(r[0]).startswith('Month'): month = int(str(r[0])[5:])
    lbl, amt = r[1], r[2]
    if isinstance(amt, str): amt = num(amt.replace('=', '').replace('\n', '').replace('\t', '').strip())
    if lbl and amt not in (None, ''):
        lbl = str(lbl)
        if lbl.startswith('Vat_'):
            m = re.match(r'Vat_(\d{4})-(\d{2})-(\d{2})', lbl); d = dt.date(int(m[1]), int(m[3]), int(m[2]))
            vat.append([d.isoformat(), r2(num(amt))])
        elif lbl.startswith('FBN'): fbn.append([month, r2(num(amt))])
        else: stLines.append([lbl, r2(num(amt))])
    if r[PD] and r[PD + 1] not in (None, ''):
        d = pay_date(str(r[PD]), month)
        if d: payments.append([d.isoformat(), r2(num(r[PD + 1]))])
payments.sort()

# ---- monthly non-order fees ----
ws = wb['All Fee']; fees = []
for r in ws.iter_rows(values_only=True):
    if r[0] and str(r[0]).startswith('Month '):
        fees.append([int(str(r[0])[6:])] + [r2(num(x)) for x in r[1:6]])

asOf = max(i[0] for i in items if i[0] is not None)
data = {'meta': {'account': ACCT, 'channel': 'Noon KSA bulky' if BULKY else 'Noon KSA courier', 'file': sys.argv[3] if len(sys.argv) > 3 else src.split('/')[-1], 'base': '2026-01-01', 'asOf': asOf,
        'status': STATUS, 'channels': CH},
        'skus': skus, 'items': items, 'statements': statements, 'stLines': stLines, 'vat': vat, 'fbn': fbn, 'payments': payments, 'fees': fees}
json.dump(data, open(out, 'w'), ensure_ascii=False, separators=(',', ':'))
print('items', len(items), 'skus', len(skus), 'titled', sum(1 for s in skus if s[2]), 'statements', len(statements), 'payments', len(payments), 'vat', len(vat), 'fbn', len(fbn), 'fees', len(fees))
print('stmt total', r2(sum(s['total'] for s in statements)), 'sheet PS lines', r2(sum(x[1] for x in stLines)), 'vat', r2(sum(x[1] for x in vat)), 'fbn', r2(sum(x[1] for x in fbn)), 'paid', r2(sum(p[1] for p in payments)))
